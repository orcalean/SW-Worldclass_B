import io

import pymupdf
import pytest
from PIL import Image

from app.ingestion import offline
from app.ingestion.loaders import UnsupportedFormat, load
from app.ingestion.structurer import structure
from app.llm.client import CallUsage, ScriptedLLM

SOP_LINES = [
    "Badge Application - Standard Work Instruction",
    "SWI ID: PLX-1955-07   Revision: 7   Effective: 2026-08-01",
    "Station: 1955L   Owner: Process Engineering",
    "1. Verify badge part number matches model and finish",
    "2. Load badge into fixture",
    "Key point: badge face down",
    "3. Locate fixture to datum holes A and B",
    "4. Apply badge with specified force, minimum hold 5 s",
    "5. Check position, level, spacing, adhesion and appearance",
]


def make_pdf(lines: list[str], with_image: bool = True, pages: int = 1) -> bytes:
    pdf = pymupdf.open()
    per = (len(lines) + pages - 1) // pages
    for i in range(pages):
        page = pdf.new_page()
        y = 72
        for ln in lines[i * per:(i + 1) * per]:
            page.insert_text((72, y), ln, fontsize=11)
            y += 20
        if with_image and i == 0:
            im = Image.new("RGB", (200, 120), (200, 220, 200))
            buf = io.BytesIO()
            im.save(buf, "PNG")
            page.insert_image(pymupdf.Rect(300, 300, 500, 420), stream=buf.getvalue())
    return pdf.tobytes()


def test_load_pdf_pages_text_images():
    doc = load("badge.pdf", make_pdf(SOP_LINES))
    assert doc.format == "pdf" and len(doc.pages) == 1
    p = doc.pages[0]
    assert "specified force" in p.text and p.render_png and len(p.images) == 1 and p.images[0].bbox


def test_unsupported_format():
    with pytest.raises(UnsupportedFormat):
        load("notes.txt", b"hello")


def test_load_docx_native():
    import docx

    d = docx.Document()
    for ln in SOP_LINES:
        d.add_paragraph(ln)
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Torque", "12 Nm"
    buf = io.BytesIO()
    d.save(buf)
    doc = load("badge.docx", buf.getvalue())
    assert "Torque | 12 Nm" in doc.pages[0].text


async def test_offline_structure_end_to_end():
    loaded = load("badge.pdf", make_pdf(SOP_LINES, pages=2))
    usage = CallUsage()
    doc = await structure(loaded, ScriptedLLM(offline.handler, "offline"), "offline", usage, pages_per_call=1)
    assert [w.seq for w in doc.work_elements] == [1, 2, 3, 4, 5]
    assert doc.meta.revision == "7" and doc.meta.swi_id == "PLX-1955-07"
    we4 = doc.work_elements[3]
    assert {(s.value, s.unit) for s in we4.specs} >= {("specified force", None), ("5", "s")}
    assert any(s.kind == "key_point" for s in doc.work_elements[1].steps)
    assert doc.work_elements[0].asset_ids == ["p1-img1"]
    assert usage.calls == 2
    assert not [w for w in doc.extraction_warnings if "captured" in w]


async def test_coverage_warning_when_content_dropped():
    loaded = load("badge.pdf", make_pdf(SOP_LINES + [f"Extra unnumbered requirement line {i} about adhesion primer"
                                                     for i in range(12)], with_image=False))

    def lossy(req):
        out = offline.handler(req)
        out.work_elements = out.work_elements[:1]
        return out

    doc = await structure(loaded, ScriptedLLM(lossy), "x", CallUsage())
    assert any("captured" in w for w in doc.extraction_warnings)


async def test_job_fails_loudly_when_every_batch_fails():
    from app.ingestion.structurer import ExtractionFailed
    from app.llm.client import LLMError

    def broken(req):
        raise LLMError("bad request: compiled grammar is too large")

    loaded = load("badge.pdf", make_pdf(SOP_LINES, pages=2))
    with pytest.raises(ExtractionFailed):
        await structure(loaded, ScriptedLLM(broken), "x", CallUsage(), pages_per_call=1)
