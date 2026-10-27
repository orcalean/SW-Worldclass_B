"""Import an OracLean core workbook (xlsx) into a versioned rubric JSON.

Expected layout (as in OracLean_Quality_Core1_Requirement_Known_Full_v1.xlsx):
  sheet 00_*  architecture: title, core question, L1-L3 roles, subcategory purposes
  sheet 01_*  one row per atomic check: Benchmark ID | Subcategory ID | Subcategory | Atomic Requirement |
              L1 | L2 | L3 | Evidence | Primary Source | Supporting Source | Notes
  sheet 03_*  evaluation logic (boundary statement)

Usage: python -m app.rubric.importer <xlsx> <out.json> --dimension Q --core 1 --version 1.0
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl

from app.contracts.rubric import Check, Core, Subcategory


def _s(v) -> str | None:
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def _sheet(wb, prefix: str):
    for ws in wb.worksheets:
        if ws.title.startswith(prefix):
            return ws
    return None


def import_core(xlsx: Path, dimension: str, number: int, version: str) -> Core:
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)

    arch = _sheet(wb, "00_")
    name, question, boundary = f"Core {number}", None, None
    level_roles: dict[str, str] = {}
    purposes: dict[str, str] = {}
    if arch is not None:
        rows = [[_s(c) for c in r] for r in arch.iter_rows(values_only=True)]
        title = rows[0][0] if rows and rows[0] else None
        if title and ":" in title:
            # "OracLean Quality — Core 1: Requirement Known — Full Draft v1"
            name = title.split(":", 1)[1].split("—")[0].strip()
        for r in rows:
            if r and r[0] and r[0].startswith(f"Core {number} asks"):
                question = r[0]
            if r and r[0] in ("L1", "L2", "L3") and len(r) > 2:
                level_roles[r[0]] = f"{r[1]}: {r[2]}"
            m = re.match(r"^(\d+\.\d+)\s", r[0] or "") if r else None
            if m and len(r) > 1 and r[1]:
                purposes[m.group(1)] = r[1]

    logic = _sheet(wb, "03_")
    if logic is not None:
        for r in logic.iter_rows(values_only=True):
            t = _s(r[0]) if r else None
            if t and "boundary" in t.lower():
                boundary = t

    full = _sheet(wb, "01_")
    if full is None:
        raise ValueError("workbook has no 01_* sheet with atomic checks")
    rows = list(full.iter_rows(values_only=True))
    header = [(_s(h) or "").lower() for h in rows[0]]

    def col(*needles: str) -> int:
        for i, h in enumerate(header):
            if all(n in h for n in needles):
                return i
        raise ValueError(f"column not found: {needles} in {header}")

    ci = {
        "id": col("benchmark id"), "sub_id": col("subcategory id"), "sub": col("subcategory"),
        "req": col("atomic"), "L1": col("l1"), "L2": col("l2"), "L3": col("l3"),
        "ev": col("evidence"), "ps": col("primary source"), "ss": col("supporting source"),
        "notes": col("notes"),
    }
    # "subcategory" also matches "subcategory id"; take the first column that is exactly the name
    ci["sub"] = next(i for i, h in enumerate(header) if h == "subcategory")

    subs: dict[str, Subcategory] = {}
    for r in rows[1:]:
        bid = _s(r[ci["id"]])
        if not bid:
            continue
        sid = _s(r[ci["sub_id"]])
        if sid not in subs:
            subs[sid] = Subcategory(id=sid, name=_s(r[ci["sub"]]) or sid, purpose=purposes.get(sid), checks=[])
        subs[sid].checks.append(Check(
            benchmark_id=bid, subcategory_id=sid, requirement=_s(r[ci["req"]]) or bid,
            criteria={lv: _s(r[ci[lv]]) or "" for lv in ("L1", "L2", "L3")},
            evidence_inputs=_s(r[ci["ev"]]), primary_source=_s(r[ci["ps"]]),
            supporting_source=_s(r[ci["ss"]]), notes=_s(r[ci["notes"]]),
        ))

    ordered = sorted(subs.values(), key=lambda s: [int(x) for x in s.id.split(".")])
    return Core(
        id=f"{dimension}.core{number}", dimension=dimension, number=number, name=name,
        question=question, boundary=boundary, level_roles=level_roles,
        subcategories=ordered, version=version,
    )


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("out", type=Path)
    ap.add_argument("--dimension", required=True)
    ap.add_argument("--core", type=int, required=True)
    ap.add_argument("--version", required=True)
    a = ap.parse_args()
    core = import_core(a.xlsx, a.dimension, a.core, a.version)
    a.out.parent.mkdir(parents=True, exist_ok=True)
    a.out.write_text(json.dumps(core.model_dump(), indent=2, ensure_ascii=False))
    print(f"{core.id}@{core.version}: {len(core.subcategories)} subcategories, {len(core.checks())} checks -> {a.out}")


if __name__ == "__main__":
    main()
