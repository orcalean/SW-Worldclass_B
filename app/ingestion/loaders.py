"""Turn any supported file into normalized pages: rendered image + text layer + embedded images.

PDF and images are native. DOCX/PPTX/XLSX are converted to PDF with LibreOffice when available (best fidelity:
layout and pictures in place); otherwise a native text+image fallback is used (no page render).
"""
from __future__ import annotations

import hashlib
import io
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from PIL import Image

MAX_EDGE = 1568  # long-edge pixels sent to the vision model
SUPPORTED = {".pdf", ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".docx", ".pptx", ".xlsx"}


@dataclass
class PageImage:
    id: str
    png: bytes
    bbox: list[float] | None = None


@dataclass
class Page:
    number: int
    text: str = ""
    render_png: bytes | None = None
    images: list[PageImage] = field(default_factory=list)


@dataclass
class LoadedDocument:
    filename: str
    format: str
    sha256: str
    pages: list[Page]
    warnings: list[str] = field(default_factory=list)


class UnsupportedFormat(ValueError):
    pass


def _to_png(data: bytes, max_edge: int = MAX_EDGE) -> bytes:
    im = Image.open(io.BytesIO(data))
    im = im.convert("RGB") if im.mode not in ("RGB", "L") else im
    im.thumbnail((max_edge, max_edge))
    buf = io.BytesIO()
    im.save(buf, "PNG", optimize=True)
    return buf.getvalue()


def _load_pdf(data: bytes) -> tuple[list[Page], list[str]]:
    pages, warnings = [], []
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        for i, p in enumerate(pdf, start=1):
            zoom = min(MAX_EDGE / max(p.rect.width, p.rect.height), 3.0)
            render = p.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).tobytes("png")
            page = Page(number=i, text=p.get_text("text").strip(), render_png=render)
            for n, img in enumerate(p.get_images(full=True)):
                xref = img[0]
                try:
                    info = pdf.extract_image(xref)
                    if info["width"] < 48 or info["height"] < 48:
                        continue  # logos/bullets/icons
                    rects = p.get_image_rects(xref)
                    bbox = [round(v, 1) for v in rects[0]] if rects else None
                    page.images.append(PageImage(id=f"p{i}-img{n + 1}", png=_to_png(info["image"]), bbox=bbox))
                except Exception as e:  # corrupt/unsupported image streams should not kill ingestion
                    warnings.append(f"page {i}: could not extract image {n + 1}: {e}")
            if not page.text:
                warnings.append(f"page {i}: no text layer (scanned?) - relying on vision")
            pages.append(page)
    return pages, warnings


def _libreoffice_to_pdf(data: bytes, suffix: str) -> bytes | None:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / f"in{suffix}"
        src.write_bytes(data)
        subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", td, str(src)],
                       check=True, capture_output=True, timeout=180)
        out = Path(td) / "in.pdf"
        return out.read_bytes() if out.exists() else None


def _native_docx(data: bytes) -> list[Page]:
    import docx

    d = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = "".join(x.text or "" for x in child.iter() if x.tag.endswith("}t")).strip()
            if t:
                lines.append(t)
        elif tag == "tbl":
            for row in child.iter():
                if row.tag.endswith("}tr"):
                    cells = []
                    for tc in row:
                        if tc.tag.endswith("}tc"):
                            cells.append(" ".join("".join(x.text or "" for x in p.iter() if x.tag.endswith("}t"))
                                                  for p in tc if p.tag.endswith("}p")).strip())
                    if any(cells):
                        lines.append(" | ".join(cells))
    page = Page(number=1, text="\n".join(lines))
    for n, rel in enumerate(r for r in d.part.rels.values() if "image" in r.reltype):
        try:
            page.images.append(PageImage(id=f"p1-img{n + 1}", png=_to_png(rel.target_part.blob)))
        except Exception:
            continue
    return [page]


def _native_pptx(data: bytes) -> list[Page]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    pages = []
    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, start=1):
        page = Page(number=i)
        texts, n = [], 0
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells))
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                n += 1
                try:
                    page.images.append(PageImage(id=f"p{i}-img{n}", png=_to_png(sh.image.blob)))
                except Exception:
                    pass
        page.text = "\n".join(texts)
        pages.append(page)
    return pages


def _native_xlsx(data: bytes) -> list[Page]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    pages = []
    for i, ws in enumerate(wb.worksheets, start=1):
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in r]
            if any(cells):
                rows.append(" | ".join(cells).strip(" |"))
        page = Page(number=i, text=f"[sheet: {ws.title}]\n" + "\n".join(rows))
        for n, img in enumerate(getattr(ws, "_images", []), start=1):
            try:
                page.images.append(PageImage(id=f"p{i}-img{n}", png=_to_png(img._data())))
            except Exception:
                pass
        pages.append(page)
    return pages


def load(filename: str, data: bytes) -> LoadedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise UnsupportedFormat(f"{suffix or 'unknown'} not supported; use one of {sorted(SUPPORTED)}")
    sha = hashlib.sha256(data).hexdigest()
    warnings: list[str] = []
    if suffix == ".pdf":
        pages, warnings = _load_pdf(data)
        fmt = "pdf"
    elif suffix in (".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"):
        png = _to_png(data)
        pages, fmt = [Page(number=1, render_png=png)], "image"
        warnings.append("image input: no text layer, relying on vision")
    else:
        fmt = suffix.lstrip(".")
        pdf = _libreoffice_to_pdf(data, suffix)
        if pdf is not None:
            pages, warnings = _load_pdf(pdf)
        else:
            pages = {"docx": _native_docx, "pptx": _native_pptx, "xlsx": _native_xlsx}[fmt](data)
            warnings.append(f"{fmt}: LibreOffice not available, using native text/image extraction "
                            "(layout not rendered; visuals may lose position context)")
    if not pages:
        raise UnsupportedFormat("document has no pages")
    return LoadedDocument(filename=filename, format=fmt, sha256=sha, pages=pages, warnings=warnings)
